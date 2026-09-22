"""会前材料解析模块：会议附件 → 文本（供纪要生成阶段注入参考，防幻觉）

支持格式：.txt/.md/.pdf/.docx/.xlsx/.csv/.pptx
并行解析，单文件失败不阻断整体。
"""
import os
import subprocess
import threading
import traceback


def parse_material(path: str, max_chars: int = 8000) -> str:
    """解析单个材料文件 → 文本（截断到 max_chars）"""
    if not os.path.exists(path):
        return f"[材料不存在: {path}]"
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".txt", ".md", ".csv"):
            with open(path, encoding="utf-8", errors="replace") as f:
                text = f.read()
        elif ext == ".pdf":
            text = _parse_pdf(path)
        elif ext == ".docx":
            text = _parse_docx(path)
        elif ext == ".xlsx":
            text = _parse_xlsx(path)
        elif ext == ".pptx":
            text = _parse_pptx(path)
        else:
            return f"[不支持的材料格式: {ext}]"
        text = text.strip()
        return text[:max_chars] if text else "[材料解析为空]"
    except Exception as e:
        return f"[材料解析失败: {path} → {e}]"


def _parse_pdf(path: str) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(path)
        return "\n".join(page.get_text() for page in doc)
    except ImportError:
        # 兜底：pdftotext（poppler）
        r = subprocess.run(["pdftotext", "-layout", path, "-"],
                           capture_output=True, timeout=60)
        return r.stdout.decode("utf-8", errors="replace")


def _parse_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"【工作表: {ws.title}】")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def _parse_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    slide_texts.append(t)
        if slide_texts:
            parts.append(f"【第{i}页】" + "\n".join(slide_texts))
    return "\n".join(parts)


SUPPORTED_EXTS = {".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx", ".pptx"}


def expand_paths(paths):
    """展开路径列表：文件直接收录；目录递归扫描其中的支持格式文件。
    目录里也排除隐藏文件/临时文件。"""
    files = []
    for p in (paths or []):
        if not p:
            continue
        if os.path.isdir(p):
            for root, dirs, fnames in os.walk(p):
                # 跳过隐藏目录
                dirs[:] = [d for d in dirs if not d.startswith(".")]
                for fn in sorted(fnames):
                    if fn.startswith(".") or fn.startswith("~$"):
                        continue
                    ext = os.path.splitext(fn)[1].lower()
                    if ext in SUPPORTED_EXTS:
                        files.append(os.path.join(root, fn))
        elif os.path.isfile(p):
            files.append(p)
    return files


def parse_materials(paths, max_chars_each: int = 8000) -> list:
    """并行解析多个材料路径（支持目录）→ [{path, name, text, ok}]"""
    paths = expand_paths(paths)
    if not paths:
        return []
    results = [None] * len(paths)

    def work(i, p):
        try:
            text = parse_material(p, max_chars_each)
            ok = not text.startswith("[材料")
            results[i] = {"path": p, "name": os.path.basename(p), "text": text, "ok": ok}
        except Exception as e:
            results[i] = {"path": p, "name": os.path.basename(p),
                          "text": f"[解析异常: {e}]", "ok": False}
            traceback.print_exc()

    threads = [threading.Thread(target=work, args=(i, p)) for i, p in enumerate(paths)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=120)
    return [r for r in results if r]


def materials_summary(materials: list) -> str:
    """材料摘要（注入纪要提示词用）"""
    if not materials:
        return ""
    ok = [m for m in materials if m.get("ok")]
    if not ok:
        return ""
    parts = []
    for m in ok:
        parts.append(f"\n【会议材料: {m['name']}】\n{m['text']}")
    return "".join(parts)


if __name__ == "__main__":
    import sys
    for p in sys.argv[1:]:
        print(f"=== {p} ===")
        print(parse_material(p, 500))
