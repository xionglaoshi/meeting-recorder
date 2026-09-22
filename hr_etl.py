#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
花名册 / 考勤 / 考核 / 日报 ETL —— 你自己的 Excel·CSV → SQLite
============================================================
用途：会议记录服务的词汇表与实体核验需要人名、部门、岗位信息；本脚本把你自己导出的
      花名册类表格重建成一个小型 SQLite 库（默认 /tmp/hr_data.db，用完即弃），
      供 vocab.json 的人名层、简称消歧、实体核验使用。

数据源: $MEETING_HR_SRC_DIR（默认 <本技能目录>/roster/）
  最省事：把花名册/通讯录丢成 roster.csv|xlsx、花名册.csv|xlsx、通讯录.csv|xlsx
          （列名含 姓名/部门/职务/入职日期 即可，中英文表头都认）
  复杂表：同目录写 hr_sources.py 列出你的表（模板见 hr_sources.example.py），
          kind 取值 assessment(考核) / attendance(考勤+花名册) / daily_report(日报) / daily_review(审核)
输出: $MEETING_HR_DB（默认 /tmp/hr_data.db）

本脚本属**可选数据源**：源目录不存在时直接以提示退出，不影响 meeting-server 的
核心录音/转写/纪要功能（相关环节自动降级）。

设计要点:
- 表头动态映射(按列名关键字匹配, 兼容各月份表头差异)
- 日期重建: 表头序列号是模板废数据(45474=2024-07-01与月份矛盾),
  日期 = sheet名取月份 + 列序号取日号
- 幂等: sources表记录sha256, 文件未变自动跳过
- 2026年考核表大量空白 → 成绩/等级存NULL(不误标0分/D)
用法: python3 hr_etl.py            # 解释器由 runtime.py 决定，见同目录 SKILL.md「运维要点」
"""
import os, re, hashlib, sqlite3, sys, datetime

# 你的表清单别写在这里（会进 Git）—— 放同目录 hr_sources.py，不进 Git
SRC_DIR = os.path.expanduser(os.environ.get("MEETING_HR_SRC_DIR")
                            or os.path.join(os.path.dirname(os.path.abspath(__file__)), "roster"))
DB_PATH = os.environ.get("MEETING_HR_DB") or "/tmp/hr_data.db"

# 默认清单只是**示例**：文件不存在就全部跳过，不会报错。
# 覆盖方式（同目录 hr_sources.py，模板 hr_sources.example.py）：
#     SRC_DIR = "~/wps/我的公司/人事导出"
#     FILES = [("2026年花名册.xlsx", "attendance", 2026), ...]
#     DAILY_REPORT_MAX_MONTH = {"26年日报汇总.xls": 7}    # 可选
FILES = [
    ("2026年花名册.xlsx", "attendance", 2026),
    ("2026年度部门月度考核表.xlsx", "assessment", 2026),
    ("26年日报汇总.xls", "daily_report", 2026),
    ("26年日报审核.xls", "daily_review", 2026),
]
DAILY_REPORT_MAX_MONTH = {}

try:  # 本地私有清单（公司表名/目录不进 Git）
    import hr_sources as _local
    if getattr(_local, "SRC_DIR", None):
        SRC_DIR = os.path.expanduser(str(_local.SRC_DIR))
    if getattr(_local, "FILES", None):
        FILES = list(_local.FILES)
    if getattr(_local, "DAILY_REPORT_MAX_MONTH", None):
        DAILY_REPORT_MAX_MONTH = dict(_local.DAILY_REPORT_MAX_MONTH)
except ImportError:
    pass

if os.environ.get("MEETING_HR_SRC_DIR"):        # 环境变量最优先，压过本地私有清单
    SRC_DIR = os.path.expanduser(os.environ["MEETING_HR_SRC_DIR"])

# ---------------- 工具 ----------------

def fullwidth_to_half(s: str) -> str:
    """全角数字→半角: 2６年 → 26年"""
    return "".join(chr(ord(c) - 0xFEE0) if '\uFF10' <= c <= '\uFF19' else c for c in s)

def extract_year(filename: str) -> int:
    m = re.search(r"(\d{1,4})\s*年", fullwidth_to_half(filename))
    if not m:
        return 0
    y = int(m.group(1))
    return 2000 + y if y < 100 else y

def extract_month(sheetname: str) -> int:
    m = re.search(r"(\d{1,2})\s*月", fullwidth_to_half(sheetname))
    return int(m.group(1)) if m else 0

def col_text(v) -> str:
    return str(v).strip() if v is not None else ""

def find_col(header_rows, keywords, exclude=()):
    """在表头行列表(每行是list[str])中, 找第一个列名包含任一关键字、且不含排除词的列。
    注意: header_rows[0] 必须是数据表头行(R2), 标题行(R1)不得传入, 否则会误匹配
    标题文本(如"各部门考核"里的"部门")。"""
    ncols = len(header_rows[0])
    for c in range(ncols):
        joined = "|".join(col_text(header_rows[r][c]) for r in range(len(header_rows)))
        if any(k in joined for k in keywords) and not any(e in joined for e in exclude):
            return c
    return -1

def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

# ---------------- 解析: 月度考核表 ----------------

def parse_assessment(conn, path, year):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    imported = 0
    for ws in wb.worksheets:
        month = extract_month(ws.title)
        if not month or "汇总" in ws.title:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 4:
            continue
        hdr = rows[1:3]  # R2大类 R3细项 (R1标题含"各部门"等字样, 必须排除)
        c_name = find_col(hdr, ["姓名"])
        c_final = find_col(hdr, ["最终成绩"])
        c_grade = find_col(hdr, ["等级"])
        c_anote = find_col(hdr, ["考勤说明"])
        c_remark = find_col(hdr, ["备注"])
        def fine(kw):
            return find_col(hdr, [kw])
        cols = {
            "attendance_10": fine("考勤执行情况"),
            "discipline_10": fine("日常劳动纪律"),
            "daily_prompt_10": fine("日报填写时效"),
            "daily_content_10": fine("日报总结内容"),
            "task_rate_20": fine("任务完成率"),
            "task_quality_20": fine("任务完成质量"),
            "bonus_20": fine("加分项"),
            "penalty_20": fine("减分项"),
        }
        if c_name < 0:
            continue
        for r in rows[3:]:
            name = col_text(r[c_name]) if c_name < len(r) else ""
            if not name or name in ("汇报人",):
                continue
            def g(idx):
                return r[idx] if 0 <= idx < len(r) else None
            ym = f"{year}-{month:02d}"
            final = g(c_final)
            grade = col_text(g(c_grade)) or None
            # 空白/占位(如实习生'/'未考核) → None; 不存0分
            def _num(v):
                if v is None:
                    return None
                if isinstance(v, (int, float)):
                    return float(v)
                s = str(v).strip()
                if not s or s in ("/", "-"):
                    return None
                try:
                    return float(s)
                except ValueError:
                    return None
            final_v = _num(final)
            conn.execute("""INSERT OR REPLACE INTO assessment_monthly
                (name, ym, dept, position, hire_date, attendance_10, discipline_10,
                 daily_prompt_10, daily_content_10, task_rate_20, task_quality_20,
                 bonus_20, penalty_20, final_score, grade, attendance_note, remark)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (name, ym,
                 col_text(g(find_col(hdr, ["部门"]))) or None,
                 col_text(g(find_col(hdr, ["岗位"]))) or None,
                 col_text(g(find_col(hdr, ["入职"]))) or None,
                 g(cols["attendance_10"]), g(cols["discipline_10"]),
                 g(cols["daily_prompt_10"]), g(cols["daily_content_10"]),
                 g(cols["task_rate_20"]), g(cols["task_quality_20"]),
                 g(cols["bonus_20"]), g(cols["penalty_20"]),
                 final_v, grade,
                 col_text(g(c_anote)) or None, col_text(g(c_remark)) or None))
            imported += 1
    wb.close()
    return imported

# ---------------- 解析: 考勤明细 ----------------

def parse_attendance(conn, path, year):
    import xlrd
    wb = xlrd.open_workbook(path)
    imported_d, imported_m, imported_e = 0, 0, 0
    for sh in wb.sheets():
        if sh.name in ("Sheet2", "花名册") or "花名册" in sh.name:
            imported_e += parse_roster(conn, sh)
            continue
        month = extract_month(sh.name)
        if not month:
            continue
        hdr = [[col_text(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(1, 3)]
        c_name = find_col(hdr, ["姓名"])
        if c_name < 0:
            continue
        c_expected = find_col(hdr, ["应出勤"])
        # 日期列 = 姓名列之后 ~ 应出勤列之前
        day_start = c_name + 1
        day_end = c_expected if c_expected > day_start else sh.ncols
        stat_cols = {
            "expected_days": find_col(hdr, ["应出勤"]),
            "overtime": find_col(hdr, ["加班"]),
            "month_rest": find_col(hdr, ["月应休"]),
            "actual_rest": find_col(hdr, ["本休"]),
            "comp_annual": find_col(hdr, ["补休"]),
            "personal_leave": find_col(hdr, ["事假"]),
            "actual_days": find_col(hdr, ["实际出勤"]),
            "pay_days": find_col(hdr, ["计薪"]),
            "remaining_comp": find_col(hdr, ["累计剩余调休"]),
            "remaining_annual": find_col(hdr, ["剩余年假"]),
            "note": find_col(hdr, ["备注"]),
            "up_miss_punish": find_col(hdr, ["上班未打卡扣分"]),
            "down_miss_punish": find_col(hdr, ["下班未打卡扣分"]),
            "late": find_col(hdr, ["迟到"]),
            "missing_report_punish": find_col(hdr, ["缺报扣分"]),
            "late_report_punish": find_col(hdr, ["补报扣分"]),
            "attendance_score": find_col(hdr, ["考勤执行情况"]),
            "daily_prompt_score": find_col(hdr, ["日报填写时效"]),
        }
        for r in range(3, sh.nrows):
            name = col_text(sh.cell_value(r, c_name))
            if not name or name in ("汇报人",):
                continue
            ym = f"{year}-{month:02d}"
            # 日粒度
            for c in range(day_start, min(day_end, sh.ncols)):
                v = col_text(sh.cell_value(r, c))
                if not v:
                    continue
                day = c - day_start + 1
                conn.execute("INSERT OR REPLACE INTO attendance_daily (name, date, mark) VALUES (?,?,?)",
                             (name, f"{ym}-{day:02d}", v))
                imported_d += 1
            # 月度汇总
            def gv(key):
                idx = stat_cols[key]
                if idx < 0 or idx >= sh.ncols:
                    return None
                v = sh.cell_value(r, idx)
                return v if isinstance(v, (int, float)) else None
            def gt(key):
                idx = stat_cols[key]
                if idx < 0:
                    return None
                v = col_text(sh.cell_value(r, idx))
                return v or None
            conn.execute("""INSERT OR REPLACE INTO attendance_monthly
                (name, ym, dept, position, hire_date, expected_days, overtime,
                 month_rest, actual_rest, comp_annual, personal_leave, actual_days,
                 pay_days, remaining_comp, remaining_annual, note, up_miss_punish,
                 down_miss_punish, late, missing_report_punish, late_report_punish,
                 attendance_score, daily_prompt_score)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (name, ym,
                 col_text(sh.cell_value(r, find_col(hdr, ["部门"]))) or None,
                 col_text(sh.cell_value(r, find_col(hdr, ["岗位"]))) or None,
                 col_text(sh.cell_value(r, find_col(hdr, ["入职"]))) or None,
                 gv("expected_days"), gv("overtime"), gv("month_rest"), gv("actual_rest"),
                 gv("comp_annual"), gv("personal_leave"), gv("actual_days"), gv("pay_days"),
                 gv("remaining_comp"), gv("remaining_annual"), gt("note"),
                 gv("up_miss_punish"), gv("down_miss_punish"), gv("late"),
                 gv("missing_report_punish"), gv("late_report_punish"),
                 gv("attendance_score"), gv("daily_prompt_score")))
            imported_m += 1
    return imported_d, imported_m, imported_e

def parse_roster(conn, sh):
    """花名册: 姓名/主体/部门/岗位/入职日期. 兼容2025(无表头, 岗位=col5, 入职=col6)与2026(有表头)"""
    n = 0
    c_name = c_dept = c_pos = c_hire = -1
    first = [col_text(sh.cell_value(0, c)) for c in range(sh.ncols)]
    if any("姓名" in v for v in first):
        hdr = [first]
        c_name = find_col(hdr, ["姓名"])
        c_dept = find_col(hdr, ["部门"])
        c_pos = find_col(hdr, ["岗位"])
        c_hire = find_col(hdr, ["入职"])
        start = 1
    else:
        # 无表头的表：按列序硬兜底（换表就改这行；能走表头识别的优先用 parse_roster_simple）
        c_name, c_dept, c_pos, c_hire = 0, 2, 5, 6
        start = 0
    for r in range(start, sh.nrows):
        name = col_text(sh.cell_value(r, c_name))
        if not name:
            continue
        hire = col_text(sh.cell_value(r, c_hire)) if c_hire >= 0 else None
        if hire and hire.replace(".", "", 1).isdigit():
            try:
                import xlrd
                hire = str(xlrd.xldate_as_datetime(float(hire), 0).date())
            except Exception:
                pass
        conn.execute("""INSERT OR REPLACE INTO employees (name, org, dept, position, hire_date)
                        VALUES (?,?,?,?,?)""",
                     (name,
                      col_text(sh.cell_value(r, 1)) or None,
                      col_text(sh.cell_value(r, c_dept)) or None,
                      col_text(sh.cell_value(r, c_pos)) or None,
                      hire))
        n += 1
    return n

# ---------------- 解析: 通用花名册（CSV / XLSX） ----------------

ROSTER_CANDIDATES = ("roster.csv", "roster.xlsx", "花名册.csv", "花名册.xlsx",
                     "通讯录.csv", "通讯录.xlsx", "员工名单.csv", "员工名单.xlsx")


def _read_text_any(path):
    """CSV 编码兜底：utf-8-sig → utf-8 → gbk → gb18030。"""
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            with open(path, encoding=enc) as fh:
                return fh.read()
        except UnicodeDecodeError:
            continue
    with open(path, encoding="utf-8", errors="replace") as fh:
        return fh.read()


def parse_roster_simple(conn, path):
    """你自己的花名册/通讯录 → employees。列名按关键字识别（中英文），认不出的列忽略。"""
    try:
        if path.lower().endswith(".csv"):
            import csv as _csv, io as _io
            rows = [r for r in _csv.reader(_io.StringIO(_read_text_any(path)))]
        else:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sh = wb[wb.sheetnames[0]]
            rows = [list(r) for r in sh.iter_rows(values_only=True)]
    except Exception as ex:
        print(f"  [错误] 花名册读取失败 {os.path.basename(path)}: {ex}")
        return 0
    rows = [r for r in rows if any(col_text(c) for c in r)]
    if not rows:
        return 0
    hdr_i = None
    for i, r in enumerate(rows[:10]):            # 表头可能不在第一行
        if re.search(r"姓名|名字|人员|name", " ".join(col_text(c) for c in r), re.I):
            hdr_i = i
            break
    if hdr_i is None:
        print(f"  [跳过] 花名册 {os.path.basename(path)}: 找不到含「姓名」的表头行")
        return 0
    hdr = [col_text(c) for c in rows[hdr_i]]

    def pick(kws):
        for j, h in enumerate(hdr):
            if any(k.lower() in h.lower() for k in kws):
                return j
        return -1

    c_name = pick(["姓名", "名字", "人员", "name"])
    c_org = pick(["主体", "公司", "单位", "org", "company"])
    c_dept = pick(["部门", "科室", "团队", "dept"])
    c_pos = pick(["职务", "职位", "岗位", "position", "title"])
    c_hire = pick(["入职", "到岗", "hire"])
    if c_name < 0:
        return 0
    n = 0
    for r in rows[hdr_i + 1:]:
        name = col_text(r[c_name]) if c_name < len(r) else ""
        if not name or name in ("姓名", "名字"):
            continue

        def g(j):
            return (col_text(r[j]) or None) if 0 <= j < len(r) else None

        conn.execute("""INSERT OR REPLACE INTO employees (name, org, dept, position, hire_date)
                        VALUES (?,?,?,?,?)""",
                     (name, g(c_org), g(c_dept), g(c_pos), g(c_hire)))
        n += 1
    return n


# ---------------- 解析: 日报汇总 ----------------

def parse_daily_report(conn, path, year, max_month=12):
    import xlrd
    import calendar as _cal
    wb = xlrd.open_workbook(path)
    imported = 0
    for sh in wb.sheets():
        month = extract_month(sh.name)
        if not month or month > max_month:
            continue
        hdr = [[col_text(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(1, 3)]
        c_name = find_col(hdr, ["姓名"])
        if c_name < 0:
            continue
        day_start = c_name + 1
        days_in_month = _cal.monthrange(year, month)[1]
        for r in range(3, sh.nrows):
            name = col_text(sh.cell_value(r, c_name))
            if not name or name in ("汇报人",):
                continue
            ym = f"{year}-{month:02d}"
            for c in range(day_start, sh.ncols):
                v = col_text(sh.cell_value(r, c))
                if not v:
                    continue
                day = c - day_start + 1
                if day > days_in_month:
                    continue
                conn.execute("INSERT OR REPLACE INTO daily_report (name, date, mark) VALUES (?,?,?)",
                             (name, f"{ym}-{day:02d}", v))
                imported += 1
    return imported

# ---------------- 解析: 日报审核 ----------------

def parse_daily_review(conn, path, year):
    import xlrd
    wb = xlrd.open_workbook(path)
    imported = 0
    for sh in wb.sheets():
        month = extract_month(sh.name)
        if not month:
            continue
        hdr = [[col_text(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(1, 3)]
        c_reviewer = find_col(hdr, ["审核人"])
        c_approve = find_col(hdr, ["审批数量", "应审人数"])
        if c_reviewer < 0:
            continue
        c_dept = find_col(hdr, ["部门"])
        day_start = max(c_reviewer, c_approve) + 1 if (c_reviewer >= 0 or c_approve >= 0) else 4
        import calendar as _cal
        days_in_month = _cal.monthrange(year, month)[1]
        for r in range(2, sh.nrows):
            reviewer = col_text(sh.cell_value(r, c_reviewer))
            if not reviewer or reviewer in ("日报审核人", "审核人"):
                continue
            dept = col_text(sh.cell_value(r, c_dept)) if c_dept >= 0 else None
            ym = f"{year}-{month:02d}"
            for c in range(day_start, sh.ncols):
                v = sh.cell_value(r, c)
                if isinstance(v, float) and v < 0:
                    day = c - day_start + 1
                    if day > days_in_month:
                        continue
                    conn.execute("INSERT OR REPLACE INTO daily_review (dept, reviewer, date, unapproved) VALUES (?,?,?,?)",
                                 (dept, reviewer, f"{ym}-{day:02d}", -v))
                    imported += 1
    return imported

# ---------------- 主流程 ----------------

def main():
    if not os.path.isdir(SRC_DIR):
        print(f"⚠ 数据源目录不存在，跳过人事库重建：\n    {SRC_DIR}\n"
              f"  （这是可选数据源，不影响会议服务的录音/转写/纪要功能；\n"
               f"   如需启用：把花名册丢进上面这个目录（roster.xlsx / 花名册.csv / 通讯录.xlsx 等），"
               f"或用 $MEETING_HR_SRC_DIR 指向你自己的导出目录）")
        return 0
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS employees (
      name TEXT PRIMARY KEY, org TEXT, dept TEXT, position TEXT, hire_date TEXT);
    CREATE TABLE IF NOT EXISTS attendance_monthly (
      name TEXT NOT NULL, ym TEXT NOT NULL, dept TEXT, position TEXT, hire_date TEXT,
      expected_days REAL, overtime REAL, month_rest REAL, actual_rest REAL,
      comp_annual REAL, personal_leave REAL, actual_days REAL, pay_days REAL,
      remaining_comp REAL, remaining_annual REAL, note TEXT,
      up_miss_punish REAL, down_miss_punish REAL, late REAL,
      missing_report_punish REAL, late_report_punish REAL,
      attendance_score REAL, daily_prompt_score REAL,
      PRIMARY KEY(name, ym));
    CREATE TABLE IF NOT EXISTS attendance_daily (
      name TEXT NOT NULL, date TEXT NOT NULL, mark TEXT, PRIMARY KEY(name, date));
    CREATE TABLE IF NOT EXISTS assessment_monthly (
      name TEXT NOT NULL, ym TEXT NOT NULL, dept TEXT, position TEXT, hire_date TEXT,
      attendance_10 REAL, discipline_10 REAL, daily_prompt_10 REAL, daily_content_10 REAL,
      task_rate_20 REAL, task_quality_20 REAL, bonus_20 REAL, penalty_20 REAL,
      final_score REAL, grade TEXT, attendance_note TEXT, remark TEXT,
      PRIMARY KEY(name, ym));
    CREATE TABLE IF NOT EXISTS daily_report (
      name TEXT NOT NULL, date TEXT NOT NULL, mark TEXT, PRIMARY KEY(name, date));
    CREATE TABLE IF NOT EXISTS daily_review (
      dept TEXT, reviewer TEXT, date TEXT, unapproved REAL,
      PRIMARY KEY(dept, reviewer, date));
    CREATE TABLE IF NOT EXISTS sources (
      file TEXT PRIMARY KEY, sha256 TEXT, rows_imported INTEGER, imported_at TEXT);
    """)
    print(f"DB: {DB_PATH}")
    for fname, ftype, year in FILES:
        path = os.path.join(SRC_DIR, fname)
        if not os.path.exists(path):
            print(f"  [跳过] 不存在: {fname}")
            continue
        digest = sha256_file(path)
        row = conn.execute("SELECT sha256 FROM sources WHERE file=?", (fname,)).fetchone()
        if row and row[0] == digest:
            print(f"  [跳过] 未变化: {fname}")
            continue
        try:
            if ftype == "assessment":
                n = parse_assessment(conn, path, year)
                print(f"  [导入] {fname}: 考核记录 {n} 条")
            elif ftype == "attendance":
                d, m, e = parse_attendance(conn, path, year)
                print(f"  [导入] {fname}: 考勤日 {d} 条, 月度汇总 {m} 条, 花名册 {e} 条")
            elif ftype == "daily_report":
                n = parse_daily_report(conn, path, year, DAILY_REPORT_MAX_MONTH.get(fname, 12))
                print(f"  [导入] {fname}: 日报日记录 {n} 条")
            elif ftype == "daily_review":
                n = parse_daily_review(conn, path, year)
                print(f"  [导入] {fname}: 审核未审记录 {n} 条")
            conn.execute("INSERT OR REPLACE INTO sources (file, sha256, rows_imported, imported_at) VALUES (?,?,?,?)",
                         (fname, digest, n if isinstance(n, int) else 0, datetime.datetime.now().isoformat(timespec="seconds")))
            conn.commit()
        except Exception as ex:
            print(f"  [错误] {fname}: {ex}")
            conn.rollback()
    # 通用花名册（CSV/XLSX）：有就顺手导入，方便没有内部人事系统的人上手
    for fn in ROSTER_CANDIDATES:
        _p = os.path.join(SRC_DIR, fn)
        if os.path.exists(_p):
            n = parse_roster_simple(conn, _p)
            if n:
                conn.commit()
                print(f"  [导入] {fn}: 花名册 {n} 人")
            break
    # 汇总统计
    for t in ["employees", "attendance_daily", "attendance_monthly", "assessment_monthly", "daily_report", "daily_review"]:
        n = conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        print(f"  表 {t}: {n} 行")
    conn.close()

if __name__ == "__main__":
    main()
